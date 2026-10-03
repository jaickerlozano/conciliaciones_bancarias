"""Informe de conciliación en Excel. La hoja principal usa fórmulas vivas (totales, saldo según
registro, conciliación y diferencia) para que el cliente pueda auditar o ajustar a mano."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.conciliaciones.informes.datos import InformeConciliacion

FORMATO_PESOS = '"$"#,##0;-"$"#,##0'
FORMATO_FECHA = "DD/MM/YYYY"

MARCA = "1D5F59"
MARCA_CLARO = "EFFAF8"
GRIS = "64748B"
LINEA = "E2E8F0"
FONDO = "F8FAFC"

F_ETIQUETA = Font(bold=True, size=9, color=MARCA)
F_TITULO = Font(bold=True, size=16, color="0F172A")
F_SUAVE = Font(size=9, color=GRIS)
F_SECCION = Font(bold=True, size=11, color="0F172A")
F_ENCABEZADO = Font(bold=True, size=9, color=MARCA)
F_NEGRITA = Font(bold=True)
R_ENCABEZADO = PatternFill("solid", fgColor=MARCA_CLARO)
R_ZEBRA = PatternFill("solid", fgColor=FONDO)
B_ENCABEZADO = Border(bottom=Side(style="medium", color=MARCA))
B_FILA = Border(bottom=Side(style="thin", color=LINEA))
B_TOTAL = Border(top=Side(style="medium", color="0F172A"))
AJUSTE = Alignment(wrap_text=True, vertical="top")

# Columnas de la hoja Conciliación
COL_FECHA, COL_NUM, COL_DETALLE, COL_REF, COL_MONTO = "B", "C", "D", "E", "F"


def _cabecera_tabla(ws: Worksheet, fila: int, titulos: list[str], desde_col: int = 2) -> None:
    for i, t in enumerate(titulos):
        c = ws.cell(fila, desde_col + i, t)
        c.font = F_ENCABEZADO
        c.fill = R_ENCABEZADO
        c.border = B_ENCABEZADO
        c.alignment = Alignment(horizontal="right" if t in ("Monto", "Cargo", "Abono") else "left")


def _hoja_conciliacion(ws: Worksheet, inf: InformeConciliacion) -> None:
    ws.title = "Conciliación"
    ws.sheet_view.showGridLines = False
    for col, ancho in {"A": 2, "B": 12, "C": 9, "D": 62, "E": 14, "F": 16}.items():
        ws.column_dimensions[col].width = ancho

    ws["B1"] = "SALDO INICIAL" if inf.es_saldo_inicial else "CONCILIACIÓN BANCARIA"
    ws["B1"].font = F_ETIQUETA
    ws["B2"] = f"{inf.periodo_texto} — {inf.comunidad}"
    ws["B2"].font = F_TITULO
    cartola = ""
    if inf.cartola_desde and inf.cartola_hasta:
        cartola = (
            f" · Cartola {'Nº ' + inf.cartola_numero + ' ' if inf.cartola_numero else ''}"
            f"({inf.cartola_desde:%d/%m/%Y} al {inf.cartola_hasta:%d/%m/%Y})"
        )
    ws["B3"] = f"{inf.banco} · Cuenta {inf.cuenta}{cartola}"
    ws["B3"].font = F_SUAVE
    estado = inf.estado_texto
    if inf.cerrada_en:
        estado += f" el {inf.cerrada_en:%d/%m/%Y %H:%M}" + (
            f" por {inf.cerrada_por}" if inf.cerrada_por else ""
        )
    elif not inf.cerrada:
        estado += " — documento preliminar, la conciliación no está cerrada"
    ws["B4"] = f"Estado: {estado}"
    ws["B4"].font = Font(size=9, color=GRIS if inf.cerrada else "BE123C", bold=not inf.cerrada)

    # --- cálculo (las celdas de totales se completan al final con referencias a cada sección)
    r = inf.resumen
    ws["B6"] = "Cálculo de la conciliación"
    ws["B6"].font = F_SECCION
    lineas = [
        (7, "", f"Saldo conciliado {inf.periodo_anterior_texto}", r.saldo_anterior),
        (8, "+", f"Ingresos de {inf.periodo_texto}", r.total_ingresos),
        (9, "−", f"Egresos de {inf.periodo_texto}", r.total_egresos),
        (10, "+", "Redondeo", r.redondeo),
        (11, "", "Saldo según registro", "=F7+F8-F9+F10"),
        (12, "+", "Cheques girados no cobrados", None),
        (13, "−", "Depósitos contabilizados no registrados en banco", None),
        (14, "+", "Movimientos no contabilizados", None),
        (15, "", "Saldo según conciliación", "=F11+F12-F13+F14"),
        (16, "", "Saldo según banco", r.saldo_banco),
        (17, "", "Diferencia (banco − conciliación)", "=F16-F15"),
    ]
    if inf.es_saldo_inicial:
        lineas[0] = (7, "", "Saldo según registro (saldo inicial)", r.saldo_registro)
    for fila, signo, etiqueta, valor in lineas:
        ws.cell(fila, 3, signo).font = F_SUAVE
        ws.cell(fila, 3).alignment = Alignment(horizontal="center")
        ws.cell(fila, 4, etiqueta)
        celda = ws.cell(fila, 6, valor)
        celda.number_format = FORMATO_PESOS
        for col in range(2, 7):
            ws.cell(fila, col).border = B_FILA
    for fila in (11, 15, 16, 17):
        ws.cell(fila, 4).font = F_NEGRITA
        ws.cell(fila, 6).font = F_NEGRITA
    for col in range(2, 7):
        ws.cell(12, col).border = Border(
            top=Side(style="medium", color="0F172A"), bottom=Side(style="thin", color=LINEA)
        )
    verde = PatternFill("solid", fgColor="ECFDF5")
    rojo = PatternFill("solid", fgColor="FFF1F2")
    ws.conditional_formatting.add(
        "D17:F17",
        CellIsRule(
            operator="equal", formula=["0"], fill=verde, font=Font(bold=True, color="047857")
        ),
    )
    ws.conditional_formatting.add(
        "D17:F17",
        CellIsRule(
            operator="notEqual", formula=["0"], fill=rojo, font=Font(bold=True, color="BE123C")
        ),
    )

    # --- secciones
    fila = 20
    total_cheques, fila = _seccion(
        ws,
        fila,
        "Cheques girados no cobrados",
        ["Fecha", "Nº", "Concepto", "Nº cheque", "Monto"],
        [
            (
                p.fecha,
                p.comprobante,
                p.detalle + (" (mes anterior)" if p.mes_anterior else ""),
                p.referencia,
                p.monto,
            )
            for p in inf.cheques
        ],
        "Todos los cheques girados fueron cobrados.",
    )
    total_depositos, fila = _seccion(
        ws,
        fila,
        "Depósitos contabilizados no registrados en banco",
        ["Fecha", "Nº", "Detalle", "Depto", "Monto"],
        [
            (
                p.fecha,
                p.comprobante,
                (p.detalle or "Gasto común") + (" (mes anterior)" if p.mes_anterior else ""),
                p.referencia,
                p.monto,
            )
            for p in inf.depositos
        ],
        "Todos los ingresos registrados aparecen en la cartola.",
    )
    total_movimientos, fila = _seccion(
        ws,
        fila,
        "Movimientos no contabilizados (abonos +, cargos −)",
        ["Fecha", "", "Descripción", "Documento", "Monto"],
        [
            (
                m.fecha,
                None,
                m.descripcion + (" (mes anterior)" if m.mes_anterior else ""),
                m.documento,
                m.monto_con_signo,
            )
            for m in inf.no_contabilizados
        ],
        "Todos los movimientos de la cartola están registrados en las planillas.",
    )
    ws["F12"] = f"={total_cheques}"
    ws["F13"] = f"={total_depositos}"
    ws["F14"] = f"={total_movimientos}"

    # firmas
    fila += 2
    for col, texto in ((2, "Preparado por"), (5, "Revisado por")):
        ws.cell(fila, col).border = Border(bottom=Side(style="thin", color="334155"))
        ws.cell(fila + 1, col, texto).font = F_SUAVE
    ws.cell(fila, 3).border = Border(bottom=Side(style="thin", color="334155"))
    ws.cell(fila, 6).border = Border(bottom=Side(style="thin", color="334155"))
    ws.cell(
        fila + 3,
        2,
        f"Generado el {inf.generado_en:%d/%m/%Y %H:%M}"
        + (f" por {inf.generado_por}" if inf.generado_por else "")
        + " · Gaudi Administraciones",
    ).font = F_SUAVE

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.left.text = f"{inf.comunidad} · {inf.periodo_texto}"


def _seccion(ws, fila, titulo, cabeceras, filas, vacio) -> tuple[str, int]:
    """Escribe una sección con su total. Devuelve (celda del total, siguiente fila libre)."""
    ws.cell(fila, 2, titulo).font = F_SECCION
    fila += 1
    _cabecera_tabla(ws, fila, cabeceras)
    fila += 1
    if not filas:
        ws.cell(fila, 2, vacio).font = Font(italic=True, size=9, color=GRIS)
        ws.cell(fila + 1, 4, "Total").font = F_NEGRITA
        ws.cell(fila + 1, 4).alignment = Alignment(horizontal="right")
        total = ws.cell(fila + 1, 6, 0)
        total.number_format = FORMATO_PESOS
        total.font = F_NEGRITA
        for col in range(2, 7):
            ws.cell(fila + 1, col).border = B_TOTAL
        return f"F{fila + 1}", fila + 4
    inicio = fila
    for i, (fecha, numero, detalle, ref, monto) in enumerate(filas):
        valores = (fecha, numero, detalle, ref or None, monto)
        for j, v in enumerate(valores):
            c = ws.cell(fila, 2 + j, v)
            c.border = B_FILA
            c.alignment = AJUSTE
            if i % 2:
                c.fill = R_ZEBRA
        ws.cell(fila, 2).number_format = FORMATO_FECHA
        ws.cell(fila, 6).number_format = FORMATO_PESOS
        fila += 1
    ws.cell(fila, 4, "Total").font = F_NEGRITA
    ws.cell(fila, 4).alignment = Alignment(horizontal="right")
    total = ws.cell(fila, 6, f"=SUM(F{inicio}:F{fila - 1})")
    total.number_format = FORMATO_PESOS
    total.font = F_NEGRITA
    for col in range(2, 7):
        ws.cell(fila, col).border = B_TOTAL
    return f"F{fila}", fila + 3


def _hoja_tabla(
    wb: Workbook,
    titulo: str,
    cabeceras: list[str],
    filas: list[tuple],
    anchos: list[int],
    columnas_fecha: set[int],
    columnas_monto: set[int],
    totales: set[int] | None = None,
) -> None:
    ws = wb.create_sheet(titulo)
    ws.sheet_view.showGridLines = False
    _cabecera_tabla(ws, 1, cabeceras, desde_col=1)
    for i, fila in enumerate(filas, start=2):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(i, j, v)
            c.border = B_FILA
            c.alignment = AJUSTE
            if j in columnas_fecha:
                c.number_format = FORMATO_FECHA
            if j in columnas_monto:
                c.number_format = FORMATO_PESOS
    ultima = len(filas) + 1
    if totales and filas:
        ws.cell(ultima + 1, 1, "Total").font = F_NEGRITA
        for j in totales:
            letra = get_column_letter(j)
            c = ws.cell(ultima + 1, j, f"=SUM({letra}2:{letra}{ultima})")
            c.number_format = FORMATO_PESOS
            c.font = F_NEGRITA
        for j in range(1, len(cabeceras) + 1):
            ws.cell(ultima + 1, j).border = B_TOTAL
    for j, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cabeceras))}{ultima}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def generar_excel(inf: InformeConciliacion) -> bytes:
    wb = Workbook()
    _hoja_conciliacion(wb.active, inf)
    if not inf.es_saldo_inicial:
        _hoja_tabla(
            wb,
            "Cruces",
            [
                "Tipo",
                "Nº",
                "Fecha libro",
                "Detalle libro",
                "Fecha banco",
                "Detalle banco",
                "Monto",
                "Tipo de cruce",
                "Confirmó",
            ],
            [
                (
                    x.partida_tipo,
                    x.comprobante,
                    x.fecha_libro,
                    x.detalle_libro,
                    x.fecha_banco,
                    x.detalle_banco,
                    x.monto,
                    x.tipo,
                    x.confirmado_por or None,
                )
                for x in inf.cruces
            ],
            [9, 7, 12, 50, 12, 40, 14, 16, 18],
            {3, 5},
            {7},
        )
        _hoja_tabla(
            wb,
            "Cartola",
            ["Fecha", "Descripción", "Documento", "Cargo", "Abono", "Estado"],
            [
                (m.fecha, m.descripcion, m.documento or None, m.cargo, m.abono, m.estado)
                for m in inf.cartola
            ],
            [12, 48, 14, 14, 14, 34],
            {1},
            {4, 5},
            totales={4, 5},
        )
    wb.properties.title = f"{inf.titulo} — {inf.comunidad}"
    wb.properties.creator = "Gaudi Administraciones"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
