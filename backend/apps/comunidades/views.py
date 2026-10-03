from openpyxl import load_workbook
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response

from apps.comunidades.models import Comunidad, CuentaBancaria
from apps.comunidades.serializers import ComunidadSerializer, CuentaBancariaSerializer
from apps.conciliaciones import servicios
from apps.conciliaciones.serializers import ConciliacionDetalleSerializer
from motor.dominio import Periodo


class ComunidadViewSet(viewsets.ModelViewSet):
    queryset = Comunidad.objects.prefetch_related("cuentas")
    serializer_class = ComunidadSerializer


class CuentaBancariaViewSet(viewsets.ModelViewSet):
    queryset = CuentaBancaria.objects.select_related("comunidad")
    serializer_class = CuentaBancariaSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        comunidad = self.request.query_params.get("comunidad")
        return qs.filter(comunidad_id=comunidad) if comunidad else qs

    @action(detail=False, methods=["post"])
    def hojas(self, request):
        """Lista las hojas de una planilla de conciliación (para elegir la de apertura)."""
        archivo = request.FILES.get("archivo")
        if archivo is None:
            raise ValidationError({"archivo": "Debe adjuntar la planilla."})
        servicios.validar_archivo(archivo, servicios.TipoArchivo.APERTURA)
        wb = load_workbook(archivo, read_only=True)
        try:
            return Response({"hojas": wb.sheetnames})
        finally:
            wb.close()

    @action(detail=True, methods=["post"])
    def apertura(self, request, pk=None):
        """Importa el estado inicial desde la última conciliación hecha a mano por el cliente."""
        archivo = request.FILES.get("archivo")
        hoja = request.data.get("hoja")
        periodo = request.data.get("periodo")
        faltan = {
            k: "Campo obligatorio."
            for k, v in {"archivo": archivo, "hoja": hoja, "periodo": periodo}.items()
            if not v
        }
        if faltan:
            raise ValidationError(faltan)
        try:
            periodo = Periodo.parse(periodo)
        except ValueError as e:
            raise ValidationError({"periodo": "Formato esperado: AAAA-MM."}) from e
        conciliacion = servicios.importar_apertura(
            self.get_object(), archivo, hoja, periodo, request.user
        )
        return Response(ConciliacionDetalleSerializer(conciliacion).data, status=201)
