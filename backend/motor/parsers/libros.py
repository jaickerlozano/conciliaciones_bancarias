"""Parser de las planillas de ingresos y egresos del cliente.

Las planillas son un listado acumulado. Cada mes termina con una fila "CIERRE MES <MES> <AÑO>"
(o "CIERRE MES DE <MES> DE <AÑO>", "CIERRE DE <MES>´<AA>"…);
el período de una partida lo define el bloque en que está, NO su fecha (hay ingresos con fecha
de junio dentro del bloque de mayo). Las filas posteriores al último cierre forman el período
abierto (el siguiente al último cierre).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from motor.dominio import PartidaLibro, Periodo, TipoPartida
from motor.texto import a_fecha, a_pesos, extraer_mes_anio, fmt_clp, normalizar

# Algunas hojas traen miles de columnas vacías; leer más allá de esto es lentísimo.
MAX_COLUMNAS = 20
FILAS_BUSQUEDA_ENCABEZADO = 15

HOJA_POR_DEFECTO = {
    TipoPartida.INGRESO: "ACUMULADO",
    TipoPartida.EGRESO: "LISTADO EGRESOS",
}

# Encabezados aceptados (normalizados) para cada campo.
ALIAS_COLUMNAS = {
    "fecha": ("FECHA",),
    "concepto": ("CONCEPTO",),
    "comprobante": ("COMPROBANTE", "COMPROB", "COMPROB.", "COMP", "COMP."),
    "torre": ("TORRE",),
    "depto": ("DEPTO", "DEPARTAMENTO"),
    "detalle": ("DETALLE", "DETALE"),
    "cheque": ("NUMERO DE CHEQUE", "N CHEQUE", "NO CHEQUE", "CHEQUE"),
    "debe": ("DEBE",),
    "haber": ("HABER",),
    "saldo": ("SALDO",),
}


class ErrorFormatoPlanilla(ValueError):
    pass


@dataclass
class LibroContable:
    tipo: TipoPartida
    bloques: dict[Periodo, list[PartidaLibro]] = field(default_factory=dict)
    periodos_cerrados: set[Periodo] = field(default_factory=set)
    advertencias: list[str] = field(default_factory=list)  # de toda la planilla
    advertencias_por_periodo: dict[Periodo, list[str]] = field(default_factory=dict)

    def advertencias_de(self, periodo: Periodo) -> list[str]:
        """Solo las advertencias sobre filas del bloque de ese período."""
        return self.advertencias_por_periodo.get(periodo, [])

    def partidas(self, periodo: Periodo) -> list[PartidaLibro]:
        return self.bloques.get(periodo, [])

    def total(self, periodo: Periodo) -> int:
        return sum(p.monto for p in self.partidas(periodo))

    @property
    def periodo_abierto(self) -> Periodo | None:
        abiertos = [p for p in self.bloques if p not in self.periodos_cerrados]
        return max(abiertos, key=lambda p: (p.anio, p.mes)) if abiertos else None


def _mapear_columnas(fila: tuple) -> dict[str, int]:
    columnas: dict[str, int] = {}
    for idx, celda in enumerate(fila):
        nombre = normalizar(celda).replace("º", "").replace("°", "")
        for campo, alias in ALIAS_COLUMNAS.items():
            if campo not in columnas and nombre in alias:
                columnas[campo] = idx
    return columnas


CONCEPTOS_PARTIDA = {"INGRESO", "INGRESOS", "EGRESO", "EGRESOS"}


def _fila_cierre(fila: tuple, columnas: dict[str, int]) -> str | None:
    """Etiqueta del cierre si la fila es un "CIERRE MES …"; si no, None.

    La etiqueta debe EMPEZAR por "CIERRE" (la celda varía según la planilla: CONCEPTO,
    COMPROBANTE o DETALLE), así una glosa como "…, CIERRE PORTÓN" no corta el bloque. Una fila
    con concepto INGRESO/EGRESO y nº de comprobante es una partida aunque su glosa empiece por
    "CIERRE"; en cambio, hay cierres con concepto EGRESO y la etiqueta en COMPROBANTE (Cinema).
    """

    def valor(campo: str):
        idx = columnas[campo]
        return fila[idx] if idx < len(fila) else None

    if normalizar(valor("concepto")) in CONCEPTOS_PARTIDA and isinstance(
        valor("comprobante"), int | float
    ):
        return None
    for celda in fila:
        if isinstance(celda, str) and normalizar(celda).startswith("CIERRE"):
            return celda
    return None


def leer_libro(ruta: str | Path, tipo: TipoPartida, hoja: str | None = None) -> LibroContable:
    hoja = hoja or HOJA_POR_DEFECTO[tipo]
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise ErrorFormatoPlanilla(
                f"No existe la hoja '{hoja}' en {Path(ruta).name}. Hojas: {wb.sheetnames}"
            )
        # Se cargan todas las filas: la columna del monto se decide mirando la planilla entera.
        filas = list(wb[hoja].iter_rows(max_col=MAX_COLUMNAS, values_only=True))
        return _procesar_filas(filas, tipo, Path(ruta).name)
    finally:
        wb.close()


def _procesar_filas(filas: list[tuple], tipo: TipoPartida, nombre_archivo: str) -> LibroContable:
    libro = LibroContable(tipo=tipo)
    columnas: dict[str, int] | None = None
    campo_monto = "debe" if tipo == TipoPartida.INGRESO else "haber"
    pendientes: list[tuple[int, PartidaLibro]] = []  # partidas del bloque en curso
    ultimo_cerrado: Periodo | None = None
    vistos: dict[int, int] = {}
    inicio_bloque = 0  # índice en libro.advertencias donde empiezan las del bloque en curso

    def asignar_advertencias(periodo: Periodo) -> None:
        nonlocal inicio_bloque
        nuevas = libro.advertencias[inicio_bloque:]
        libro.advertencias_por_periodo.setdefault(periodo, []).extend(nuevas)
        inicio_bloque = len(libro.advertencias)

    for nro_fila, fila in enumerate(filas, start=1):
        if columnas is None:
            if nro_fila > FILAS_BUSQUEDA_ENCABEZADO:
                raise ErrorFormatoPlanilla(
                    f"{nombre_archivo}: no se encontró la fila de encabezados (FECHA, CONCEPTO…)."
                )
            candidatas = _mapear_columnas(fila)
            claves = candidatas.keys()
            if {"fecha", "concepto", "comprobante"} <= claves and ({"debe", "haber"} & claves):
                columnas = candidatas
                campo_monto = _columna_monto(filas[nro_fila:], columnas, tipo)
            continue

        etiqueta = _fila_cierre(fila, columnas)
        if etiqueta is not None:
            mes_anio = extraer_mes_anio(etiqueta)
            if mes_anio is None:
                libro.advertencias.append(
                    f"Fila {nro_fila}: no se pudo leer el mes del cierre '{etiqueta.strip()}'; "
                    "sus partidas quedan sin período."
                )
                pendientes = []
                inicio_bloque = len(libro.advertencias)
                continue
            periodo = Periodo(*mes_anio)
            if periodo in libro.periodos_cerrados:
                libro.advertencias.append(
                    f"Fila {nro_fila}: el período {periodo} aparece cerrado dos veces."
                )
            libro.bloques.setdefault(periodo, []).extend(p for _, p in pendientes)
            libro.periodos_cerrados.add(periodo)
            _validar_total_cierre(libro, periodo, fila, columnas, nro_fila)
            asignar_advertencias(periodo)
            pendientes = []
            ultimo_cerrado = periodo
            continue

        partida = _leer_partida(fila, columnas, tipo, campo_monto, nro_fila, libro)
        if partida is not None:
            if partida.comprobante is not None:
                if partida.comprobante in vistos:
                    libro.advertencias.append(
                        f"Fila {nro_fila}: comprobante {partida.comprobante} duplicado "
                        f"(ya aparece en la fila {vistos[partida.comprobante]})."
                    )
                vistos[partida.comprobante] = nro_fila
            pendientes.append((nro_fila, partida))

    if columnas is None:
        raise ErrorFormatoPlanilla(f"{nombre_archivo}: la hoja está vacía o sin encabezados.")

    if pendientes:
        if ultimo_cerrado is None:
            libro.advertencias.append(
                "La planilla no tiene ninguna fila 'CIERRE MES'; no se pueden asignar períodos."
            )
        else:
            libro.bloques.setdefault(ultimo_cerrado.siguiente(), []).extend(
                p for _, p in pendientes
            )
            asignar_advertencias(ultimo_cerrado.siguiente())

    for periodo, partidas in libro.bloques.items():
        _validar_fechas(libro, periodo, partidas)
    return libro


def _leer_partida(
    fila: tuple,
    columnas: dict[str, int],
    tipo: TipoPartida,
    campo_monto: str,
    nro_fila: int,
    libro: LibroContable,
) -> PartidaLibro | None:
    def celda(campo: str):
        idx = columnas.get(campo)
        return fila[idx] if idx is not None and idx < len(fila) else None

    if normalizar(celda("concepto")) not in (tipo.value, tipo.value + "S"):
        return None
    comprobante = celda("comprobante")
    try:
        comprobante = int(comprobante) if comprobante is not None else None
    except (TypeError, ValueError):
        comprobante = None

    # La columna del monto es la misma para toda la planilla (ver `_columna_monto`). Un monto
    # solo en la columna contraria suele ser un reverso o anulación: no se suma como partida.
    otro_campo = "haber" if campo_monto == "debe" else "debe"
    bruto, otro = celda(campo_monto), celda(otro_campo)
    otro_con_monto = _es_monto(otro)
    if bruto in (None, "", 0) and otro_con_monto:
        libro.advertencias.append(
            f"Fila {nro_fila}: comprobante {comprobante} trae el monto en {otro_campo.upper()} "
            f"({fmt_clp(a_pesos(otro)[0])}) y esta planilla usa {campo_monto.upper()}; se omite. "
            "Revise si es un reverso o un error de digitación."
        )
        return None
    if otro_con_monto and isinstance(bruto, int | float):
        libro.advertencias.append(
            f"Fila {nro_fila}: comprobante {comprobante} tiene monto en DEBE y en HABER; "
            f"se usa {campo_monto.upper()} ({fmt_clp(a_pesos(bruto)[0])}). Revise la planilla."
        )
    if bruto in (None, "") or not isinstance(bruto, int | float):
        if bruto not in (None, ""):
            libro.advertencias.append(
                f"Fila {nro_fila}: monto no numérico '{bruto}' (comprobante {comprobante}); "
                "se omite."
            )
        else:
            libro.advertencias.append(
                f"Fila {nro_fila}: comprobante {comprobante} sin monto; se omite."
            )
        return None
    monto, tenia_decimales = a_pesos(bruto)
    if monto == 0:
        return None  # comprobantes anulados ("NULO") vienen con monto 0
    if tenia_decimales:
        libro.advertencias.append(
            f"Fila {nro_fila}: comprobante {comprobante} tiene decimales ({bruto}); "
            f"se redondea a {fmt_clp(monto)}."
        )

    fecha = a_fecha(celda("fecha"))
    if fecha is None:
        libro.advertencias.append(f"Fila {nro_fila}: comprobante {comprobante} sin fecha válida.")

    cheque = celda("cheque")
    if isinstance(cheque, float) and cheque.is_integer():
        cheque = int(cheque)
    depto = _texto_entero(celda("depto"))
    torre = _texto_entero(celda("torre"))
    if torre and depto:
        depto = f"{torre}-{depto}"  # en edificios con torres el mismo depto existe en ambas

    return PartidaLibro(
        tipo=tipo,
        comprobante=comprobante,
        fecha=fecha,
        monto=monto,
        glosa=str(celda("detalle") or "").strip(),
        depto=depto,
        cheque=str(cheque or "").strip(),
    )


def _es_monto(valor) -> bool:
    return isinstance(valor, int | float) and not isinstance(valor, bool) and valor != 0


def _columna_monto(filas: list[tuple], columnas: dict[str, int], tipo: TipoPartida) -> str:
    """Columna del monto para toda la planilla: la que tiene más partidas de ese tipo con monto.

    Lo normal es ingresos en DEBE y egresos en HABER, pero hay planillas al revés (ej. ingresos
    de General Córdova, en HABER). En empate gana la columna habitual.
    """
    esperada = "debe" if tipo == TipoPartida.INGRESO else "haber"
    otra = "haber" if esperada == "debe" else "debe"
    conceptos = (tipo.value, tipo.value + "S")
    i_concepto = columnas["concepto"]
    cuenta = {esperada: 0, otra: 0}
    for fila in filas:
        if i_concepto >= len(fila) or normalizar(fila[i_concepto]) not in conceptos:
            continue
        for campo in cuenta:
            idx = columnas.get(campo)
            if idx is not None and idx < len(fila) and _es_monto(fila[idx]):
                cuenta[campo] += 1
    return otra if cuenta[otra] > cuenta[esperada] else esperada


def _texto_entero(valor) -> str:
    """Texto de una celda; los números enteros leídos como float (61.0) pierden el decimal."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor or "").strip()


def _validar_total_cierre(
    libro: LibroContable, periodo: Periodo, fila: tuple, columnas: dict[str, int], nro_fila: int
) -> None:
    """Si la fila de cierre trae el total acumulado del mes (columna SALDO), se compara."""
    idx = columnas.get("saldo")
    total_fila = fila[idx] if idx is not None and idx < len(fila) else None
    if isinstance(total_fila, int | float) and total_fila:
        calculado = libro.total(periodo)
        informado, _ = a_pesos(total_fila)
        if calculado != informado:
            libro.advertencias.append(
                f"Fila {nro_fila}: el cierre de {periodo} informa {fmt_clp(informado)} pero las "
                f"partidas suman {fmt_clp(calculado)}."
            )


def _validar_fechas(libro: LibroContable, periodo: Periodo, partidas: list[PartidaLibro]) -> None:
    for p in partidas:
        if p.fecha and abs(p.fecha.year - periodo.anio) > 1:
            aviso = f"Comprobante {p.comprobante} ({periodo}): fecha sospechosa {p.fecha:%d/%m/%Y}."
            libro.advertencias.append(aviso)
            libro.advertencias_por_periodo.setdefault(periodo, []).append(aviso)
