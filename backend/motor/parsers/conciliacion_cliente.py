"""Lee una hoja de la planilla "CONCILIACIÓN MENSUAL" del cliente.

Sirve para dos cosas:
- Obtener el estado de apertura (pendientes y saldo) al empezar a usar el sistema con una
  comunidad, a partir de la última conciliación hecha a mano.
- Comparar los resultados del motor contra las conciliaciones históricas (tests).

Estructura esperada (una hoja por mes):
    "Saldo según registro"                         -> monto en la misma fila
    FECHA | (tipo) | Nº | CONCEPTO | ... | Nº CHEQUE | MONTO | TOTALES   (fila de encabezados)
    "Cheques girados no cobrados"                  -> filas de egresos pendientes
    "Depósitos contabilizados no registrados..."   -> filas de ingresos pendientes
    "MOVIMIENTOS NO CONTABILIZADOS"                -> filas de movimientos del banco
    "Saldo según conciliación" / "Saldo según banco"
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

from motor.dominio import EstadoApertura, MovimientoBancario, PartidaLibro, TipoPartida
from motor.texto import a_fecha, a_pesos, normalizar

MAX_COLUMNAS = 20


class ErrorFormatoConciliacion(ValueError):
    pass


@dataclass
class ConciliacionCliente:
    hoja: str
    saldo_anterior: int | None = None
    total_ingresos: int | None = None
    total_egresos: int | None = None
    redondeo: int = 0
    saldo_registro: int | None = None
    saldo_banco: int | None = None
    cheques_pendientes: list[PartidaLibro] = field(default_factory=list)
    depositos_pendientes: list[PartidaLibro] = field(default_factory=list)
    movimientos_no_contabilizados: list[MovimientoBancario] = field(default_factory=list)

    def como_apertura(self) -> EstadoApertura:
        if self.saldo_registro is None:
            raise ErrorFormatoConciliacion(
                f"La hoja '{self.hoja}' no tiene 'Saldo según registro'."
            )
        return EstadoApertura(
            saldo_registro=self.saldo_registro,
            saldo_banco=self.saldo_banco,
            cheques_pendientes=self.cheques_pendientes,
            depositos_pendientes=self.depositos_pendientes,
            movimientos_no_contabilizados=self.movimientos_no_contabilizados,
        )


_SECCIONES = {
    "CHEQUES GIRADOS NO COBRADOS": "cheques",
    "DEPOSITOS CONTABILIZADOS NO REGISTRADOS EN BANCO": "depositos",
    "MOVIMIENTOS NO CONTABILIZADOS": "movimientos",
}


def leer_conciliacion_cliente(ruta: str | Path, hoja: str) -> ConciliacionCliente:
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise ErrorFormatoConciliacion(f"No existe la hoja '{hoja}' en {Path(ruta).name}.")
        filas = list(wb[hoja].iter_rows(max_col=MAX_COLUMNAS, values_only=True))
    finally:
        wb.close()
    return _procesar(filas, hoja)


def _ultimo_numero(fila: tuple) -> int | None:
    for celda in reversed(fila):
        if isinstance(celda, int | float):
            return a_pesos(celda)[0]
    return None


def _procesar(filas: list[tuple], hoja: str) -> ConciliacionCliente:
    resultado = ConciliacionCliente(hoja=hoja)
    col: dict[str, int] = {}
    seccion: str | None = None

    for fila in filas:
        textos = [normalizar(c) for c in fila if isinstance(c, str)]
        linea = " ".join(textos)

        if "SALDO CONCILIADO MES ANTERIOR" in linea:
            resultado.saldo_anterior = _ultimo_numero(fila)
        elif linea.startswith("INGRESOS DE"):
            resultado.total_ingresos = _ultimo_numero(fila)
        elif linea.startswith("EGRESOS DE"):
            resultado.total_egresos = _ultimo_numero(fila)
        elif linea.startswith("REDONDEO"):
            resultado.redondeo = _ultimo_numero(fila) or 0
        elif "SALDO SEGUN REGISTRO" in linea:
            resultado.saldo_registro = _ultimo_numero(fila)
        elif "SALDO SEGUN BANCO" in linea:
            resultado.saldo_banco = _ultimo_numero(fila)
        elif "SALDO SEGUN CONCILIACION" in linea:
            seccion = None
        elif not col and "FECHA" in textos and "MONTO" in textos:
            for idx, c in enumerate(fila):
                nombre = normalizar(c).replace("º", "").replace("°", "")
                if nombre == "FECHA":
                    col["fecha"] = idx
                elif nombre in ("N", "NO"):
                    col["comprobante"] = idx
                elif nombre == "CONCEPTO":
                    col["concepto"] = idx
                elif nombre in ("N CHEQUE", "NO CHEQUE"):
                    col["cheque"] = idx
                elif nombre == "MONTO":
                    col["monto"] = idx
        elif linea in _SECCIONES:
            seccion = _SECCIONES[linea]
        elif seccion and col:
            _leer_item(fila, col, seccion, resultado)

    if not col:
        raise ErrorFormatoConciliacion(f"La hoja '{hoja}' no tiene la fila de encabezados.")
    return resultado


def _leer_item(fila: tuple, col: dict[str, int], seccion: str, r: ConciliacionCliente) -> None:
    def celda(campo: str):
        idx = col.get(campo)
        return fila[idx] if idx is not None and idx < len(fila) else None

    bruto = celda("monto")
    if not isinstance(bruto, int | float):
        return
    monto, _ = a_pesos(bruto)
    fecha = a_fecha(celda("fecha"))
    idx_fecha = col["fecha"]
    comprobante = celda("comprobante")
    comprobante = int(comprobante) if isinstance(comprobante, int | float) else None

    if seccion == "movimientos":
        r.movimientos_no_contabilizados.append(
            MovimientoBancario(
                fecha=fecha or date.min,
                descripcion=" ".join(
                    str(c).strip() for c in fila[idx_fecha + 1 : col["monto"]] if c is not None
                ),
                monto=abs(monto),
                es_cargo=monto < 0,
            )
        )
        return

    tipo = TipoPartida.EGRESO if seccion == "cheques" else TipoPartida.INGRESO
    concepto = celda("concepto")
    if tipo == TipoPartida.EGRESO:
        cheque = celda("cheque")
        partida = PartidaLibro(
            tipo=tipo,
            comprobante=comprobante,
            fecha=fecha,
            monto=monto,
            glosa=str(concepto or "").strip(),
            cheque=str(int(cheque) if isinstance(cheque, float) else cheque or "").strip(),
        )
        r.cheques_pendientes.append(partida)
    else:
        # en depósitos, la columna CONCEPTO trae el depto y la siguiente el detalle ("G.C")
        idx_concepto = col["concepto"]
        detalle = fila[idx_concepto + 1] if idx_concepto + 1 < len(fila) else None
        depto = int(concepto) if isinstance(concepto, float) else concepto
        partida = PartidaLibro(
            tipo=tipo,
            comprobante=comprobante,
            fecha=fecha,
            monto=monto,
            glosa=str(detalle or "").strip(),
            depto=str(depto or "").strip(),
        )
        r.depositos_pendientes.append(partida)
