"""Plantilla estándar de cartola (Excel).

Formato de respaldo para cualquier banco cuyo archivo el sistema aún no lee (o PDFs escaneados):
el usuario copia los movimientos a esta plantilla. Estructura:

    A1  CARTOLA ESTÁNDAR
    Banco | Santander          (filas etiqueta | valor, en cualquier orden)
    Cuenta | 0-000-03-81745-8
    Desde | 01/01/2026
    Hasta | 31/01/2026
    Saldo inicial | 16356758
    Saldo final | 11015854
    Fecha | Descripción | Nº Documento | Cargo | Abono      (encabezado de la tabla)
    ...una fila por movimiento...
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from motor.dominio import Cartola, MovimientoBancario
from motor.parsers.cartolas.base import ErrorCartola, ParserCartola, registrar
from motor.texto import a_fecha, a_pesos, normalizar

MARCA = "CARTOLA ESTANDAR"
CAMPOS = ("Banco", "Cuenta", "Desde", "Hasta", "Saldo inicial", "Saldo final")
COLUMNAS = ("Fecha", "Descripción", "Nº Documento", "Cargo", "Abono")
MAX_COLUMNAS = 10


def _monto(valor, fila: int, campo: str) -> int | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, str):
        valor = valor.replace("$", "").replace(".", "").replace(",", ".").strip()
    try:
        monto, _ = a_pesos(valor)
    except Exception as e:
        raise ErrorCartola(
            f"Plantilla, fila {fila}: '{valor}' no es un monto válido ({campo})."
        ) from e
    return monto


@registrar
class PlantillaEstandar(ParserCartola):
    banco = "Plantilla estándar"
    formato = "excel-plantilla"
    extensiones = (".xlsx",)

    @classmethod
    def puede_leer(cls, ruta: Path) -> bool:
        try:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        except Exception:
            return False
        try:
            primera = next(wb.active.iter_rows(max_row=1, max_col=3, values_only=True), ())
            return any(MARCA in normalizar(c) for c in primera if c)
        finally:
            wb.close()

    def leer(self, ruta: Path) -> Cartola:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        try:
            filas = list(wb.active.iter_rows(max_col=MAX_COLUMNAS, values_only=True))
        finally:
            wb.close()

        datos: dict[str, object] = {}
        columnas: dict[str, int] | None = None
        movimientos: list[MovimientoBancario] = []
        for nro, fila in enumerate(filas, start=1):
            if columnas is None:
                etiqueta = normalizar(fila[0]) if fila else ""
                for campo in CAMPOS:
                    if etiqueta == normalizar(campo):
                        datos[campo] = fila[1] if len(fila) > 1 else None
                encabezados = [normalizar(c) for c in fila]
                if "FECHA" in encabezados and "CARGO" in encabezados and "ABONO" in encabezados:
                    columnas = {normalizar(c): i for i, c in enumerate(fila) if c}
                continue
            if all(c in (None, "") for c in fila):
                continue
            movimientos.append(self._movimiento(fila, columnas, nro))

        faltan = [c for c in CAMPOS if datos.get(c) in (None, "")]
        if faltan:
            raise ErrorCartola(f"A la plantilla de cartola le falta: {', '.join(faltan)}.")
        if columnas is None:
            raise ErrorCartola("La plantilla no tiene la fila de encabezados Fecha | … | Abono.")
        desde, hasta = a_fecha(datos["Desde"]), a_fecha(datos["Hasta"])
        if desde is None or hasta is None:
            raise ErrorCartola("En la plantilla, Desde y Hasta deben ser fechas.")
        return Cartola(
            banco=str(datos["Banco"]).strip(),
            cuenta=str(datos["Cuenta"]).strip(),
            numero="",
            desde=desde,
            hasta=hasta,
            saldo_inicial=_monto(datos["Saldo inicial"], 0, "saldo inicial"),
            saldo_final=_monto(datos["Saldo final"], 0, "saldo final"),
            movimientos=movimientos,
        )

    @staticmethod
    def _movimiento(fila: tuple, columnas: dict[str, int], nro: int) -> MovimientoBancario:
        def celda(nombre: str):
            idx = columnas.get(nombre)
            return fila[idx] if idx is not None and idx < len(fila) else None

        fecha = a_fecha(celda("FECHA"))
        if fecha is None:
            raise ErrorCartola(f"Plantilla, fila {nro}: la fecha no es válida.")
        cargo = _monto(celda("CARGO"), nro, "cargo")
        abono = _monto(celda("ABONO"), nro, "abono")
        if (cargo is None) == (abono is None):
            raise ErrorCartola(
                f"Plantilla, fila {nro}: indique el monto en Cargo o en Abono (uno solo)."
            )
        documento = celda("NO DOCUMENTO") or celda("N DOCUMENTO") or celda("Nº DOCUMENTO")
        if isinstance(documento, float) and documento.is_integer():
            documento = int(documento)
        return MovimientoBancario(
            fecha=fecha,
            descripcion=str(celda("DESCRIPCION") or "").strip(),
            monto=cargo if cargo is not None else abono,
            es_cargo=cargo is not None,
            documento=str(documento or "").strip(),
        )


def escribir_plantilla(ruta: str | Path, cartola: Cartola | None = None) -> None:
    """Crea la plantilla (vacía o con los datos de `cartola`)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cartola"
    negrita = Font(bold=True)
    ws["A1"] = "CARTOLA ESTÁNDAR"
    ws["A1"].font = Font(bold=True, size=14)
    ws["C1"] = (
        "Complete los datos y una fila por movimiento. Montos sin signo: en Cargo o en Abono."
    )
    ws["C1"].font = Font(italic=True, color="666666")
    valores = (
        (cartola.banco, cartola.cuenta, cartola.desde, cartola.hasta,
         cartola.saldo_inicial, cartola.saldo_final)
        if cartola else (None,) * len(CAMPOS)
    )  # fmt: skip
    for i, (campo, valor) in enumerate(zip(CAMPOS, valores, strict=True), start=3):
        ws.cell(i, 1, campo).font = negrita
        celda = ws.cell(i, 2, valor)
        if campo in ("Desde", "Hasta"):
            celda.number_format = "DD/MM/YYYY"
        elif campo.startswith("Saldo"):
            celda.number_format = "#,##0"

    fila_enc = 3 + len(CAMPOS) + 1
    relleno = PatternFill("solid", fgColor="DDEBF7")
    for j, nombre in enumerate(COLUMNAS, start=1):
        c = ws.cell(fila_enc, j, nombre)
        c.font = negrita
        c.fill = relleno
        c.alignment = Alignment(horizontal="center")

    for i, m in enumerate(cartola.movimientos if cartola else [], start=fila_enc + 1):
        ws.cell(i, 1, m.fecha).number_format = "DD/MM/YYYY"
        ws.cell(i, 2, m.descripcion)
        ws.cell(i, 3, m.documento or None)
        ws.cell(i, 4 if m.es_cargo else 5, m.monto).number_format = "#,##0"

    fechas = DataValidation(type="date", allow_blank=True)
    ws.add_data_validation(fechas)
    fechas.add(f"A{fila_enc + 1}:A2000")
    for col, ancho in zip("ABCDE", (14, 45, 16, 14, 14), strict=True):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = ws.cell(fila_enc + 1, 1)
    wb.save(ruta)
