"""Flujo completo por API con los archivos reales del piloto (Cinema, mayo 2026)."""

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

pytestmark = [pytest.mark.django_db, pytest.mark.datos_reales]


def _archivo(datos, nombre):
    return SimpleUploadedFile(nombre, (datos / nombre).read_bytes())


def test_flujo_mensual_completo(api, datos):
    comunidad = api.post("/api/comunidades/", {"nombre": "Comunidad Edificio Cinema"}).json()
    cuenta = api.post(
        "/api/cuentas/",
        {"comunidad": comunidad["id"], "banco": "santander", "numero": "0-000-03-81745-8"},
    ).json()

    planilla = "CONCILIACIÓN  MENSUAL CINEMA.xlsm"
    hojas = api.post(
        "/api/cuentas/hojas/", {"archivo": _archivo(datos, planilla)}, format="multipart"
    ).json()["hojas"]
    assert "ABRIL´26" in hojas

    r = api.post(
        f"/api/cuentas/{cuenta['id']}/apertura/",
        {"archivo": _archivo(datos, planilla), "hoja": "ABRIL´26", "periodo": "2026-04"},
        format="multipart",
    )
    assert r.status_code == 201, r.json()
    abril = r.json()
    assert abril["estado"] == "importada"
    assert abril["resumen"]["diferencia"] == 0
    assert len(abril["cheques_pendientes"]) == 12

    mayo = api.post("/api/conciliaciones/", {"cuenta": cuenta["id"], "periodo": "2026-05"})
    url = f"/api/conciliaciones/{mayo.json()['id']}"
    for tipo, nombre in [
        ("ingresos", "listado ingresos CINEMA2.xlsx"),
        ("egresos", "emitir egresos CINEMA.xlsm"),
        ("cartola", "cartola_mayo_2026.pdf"),
    ]:
        r = api.post(
            f"{url}/archivos/", {"tipo": tipo, "archivo": _archivo(datos, nombre)},
            format="multipart",
        )  # fmt: skip
        assert r.status_code == 201, r.json()

    r = api.post(f"{url}/procesar/")
    assert r.status_code == 200, r.json()
    detalle = r.json()
    resumen = detalle["resumen"]
    assert resumen["saldo_registro"] == 1_677_679
    assert resumen["total_cheques_pendientes"] == 6_131_348
    assert resumen["total_depositos_pendientes"] == 2_644_177
    assert resumen["saldo_banco"] == 5_164_850
    assert resumen["diferencia"] == 0
    assert resumen["cruces_por_revisar"] == 7

    r = api.post(f"{url}/cerrar/")
    assert r.status_code == 400
    assert "7 cruces por revisar" in r.json()["detail"]

    for cruce in detalle["cruces"]:
        if cruce["requiere_revision"]:
            assert api.post(f"{url}/cruces/{cruce['id']}/confirmar/").status_code == 200
    assert api.post(f"{url}/cerrar/").status_code == 200

    junio = api.post("/api/conciliaciones/", {"cuenta": cuenta["id"], "periodo": "2026-06"})
    assert junio.status_code == 201


def test_rechaza_cartola_de_otra_cuenta(api, datos, apertura):
    """La cartola de Cinema no puede usarse para conciliar otra cuenta."""
    mayo = api.post("/api/conciliaciones/", {"cuenta": apertura.cuenta_id, "periodo": "2026-05"})
    url = f"/api/conciliaciones/{mayo.json()['id']}"
    for tipo, nombre in [
        ("ingresos", "listado ingresos CINEMA2.xlsx"),
        ("egresos", "emitir egresos CINEMA.xlsm"),
        ("cartola", "cartola_mayo_2026.pdf"),
    ]:
        api.post(f"{url}/archivos/", {"tipo": tipo, "archivo": _archivo(datos, nombre)},
                 format="multipart")  # fmt: skip
    r = api.post(f"{url}/procesar/")
    assert r.status_code == 400
    assert "0-000-03-81745-8" in r.json()["detail"]


def test_informes_de_mayo_real(api, datos):
    """El Excel de mayo, evaluando sus fórmulas, da los mismos totales que la conciliación."""
    from io import BytesIO

    from openpyxl import load_workbook

    from tests.test_api.test_informes import evaluar

    test_flujo_mensual_completo(api, datos)
    mayo = api.get("/api/conciliaciones/?cuenta=" + str(
        api.get("/api/cuentas/").json()[0]["id"])).json()  # fmt: skip
    mayo = next(c for c in mayo if c["periodo"] == "2026-05")
    ws = load_workbook(BytesIO(api.get(f"/api/conciliaciones/{mayo['id']}/excel/").content)).active
    assert evaluar(ws, "F11") == 1_677_679
    assert evaluar(ws, "F12") == 6_131_348
    assert evaluar(ws, "F13") == 2_644_177
    assert evaluar(ws, "F16") == 5_164_850
    assert evaluar(ws, "F17") == 0
    assert "Cartola Nº 306 (30/04/2026 al 29/05/2026)" in ws["B3"].value
    pdf = api.get(f"/api/conciliaciones/{mayo['id']}/pdf/")
    assert pdf.content.startswith(b"%PDF")
