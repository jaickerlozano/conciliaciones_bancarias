"""Informes PDF y Excel."""

import re
from io import BytesIO

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db


def evaluar(ws, celda: str) -> int:
    """Evalúa las fórmulas simples del informe (=F1+F2-F3, =SUM(F5:F9), =F10) como Excel."""
    valor = ws[celda].value
    if not isinstance(valor, str) or not valor.startswith("="):
        return valor or 0
    expr = valor[1:]
    if m := re.fullmatch(r"SUM\(([A-Z])(\d+):([A-Z])(\d+)\)", expr):
        col, desde, _, hasta = m.groups()
        return sum(evaluar(ws, f"{col}{i}") for i in range(int(desde), int(hasta) + 1))
    total, signo = 0, 1
    for token in re.findall(r"[+-]|[A-Z]+\d+", expr):
        if token in "+-":
            signo = 1 if token == "+" else -1
        else:
            total += signo * evaluar(ws, token)
    return total


def test_pdf(api, mayo_procesado):
    r = api.get(f"/api/conciliaciones/{mayo_procesado.id}/pdf/")
    assert r.status_code == 200
    assert r["Content-Type"] == "application/pdf"
    assert r["Content-Disposition"] == (
        'attachment; filename="Conciliacion_Edificio_Prueba_2026-05.pdf"'
    )
    assert r.content.startswith(b"%PDF")


def test_excel_con_formulas_que_cuadran(api, mayo_procesado):
    r = api.get(f"/api/conciliaciones/{mayo_procesado.id}/excel/")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Conciliación", "Cruces", "Cartola"]
    ws = wb["Conciliación"]
    assert ws["F11"].value == "=F7+F8-F9+F10"
    assert evaluar(ws, "F11") == 1_050_000  # saldo según registro
    assert evaluar(ws, "F13") == 50_000  # depósito sin cruzar
    assert evaluar(ws, "F14") == 50_000  # abono sin cruzar
    assert evaluar(ws, "F16") == 1_050_000  # banco
    assert evaluar(ws, "F17") == 0  # diferencia
    assert "no está cerrada" in ws["B4"].value
    assert wb["Cartola"]["F2"].value == "No contabilizado"


def test_sin_procesar_no_hay_informe(api, apertura):
    mayo = api.post("/api/conciliaciones/", {"cuenta": apertura.cuenta_id, "periodo": "2026-05"})
    r = api.get(f"/api/conciliaciones/{mayo.json()['id']}/pdf/")
    assert r.status_code == 400
    assert "después de procesar" in r.json()["detail"]


def test_saldo_inicial_tiene_informe(api, apertura):
    r = api.get(f"/api/conciliaciones/{apertura.id}/excel/")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Conciliación"]
    assert evaluar(wb.active, "F17") == 0
